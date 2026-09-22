"""Build an editable user briefing plus same-layout PDF/PNG previews.

Dependencies: python-pptx, Pillow. Run from any working directory.
Preview images are rendered from the same layout commands, not by PowerPoint.
"""
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / 'forcing_products'
OUT.mkdir(parents=True, exist_ok=True)
W, H, SCALE = 13.333, 7.5, 120
NAVY, TEAL, BLUE, GOLD = '132E40', '087F8C', '3975B7', 'DDA530'
BG, INK, GRAY, LIGHT = 'F5F7F8', '183445', '526674', 'E3EBEF'
WHITE, RED = 'FFFFFF', 'B95140'
FONT = '/home/mpan/local/miniforge3/fonts/DejaVuSans.ttf'
BOLD = '/home/mpan/local/miniforge3/fonts/DejaVuSans-Bold.ttf'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Informed forcing for CONUS-wide NWM operations'
prs.core_properties.subject = 'User briefing: design, sources, NRT and retrospective products'
prs.core_properties.author = 'Hydro Ops'
slides, notes = [], []

SOURCES = {
 'NLDAS': 'https://ldas.gsfc.nasa.gov/nldas',
 'NLDAS forcing': 'https://ldas.gsfc.nasa.gov/nldas/v2/forcing',
 'HRRR': 'https://emc.ncep.noaa.gov/emc/pages/numerical_forecast_systems/hrrr.php',
 'MRMS': 'https://www.nssl.noaa.gov/projects/mrms/operational/tables.php',
 'MRMS versions': 'https://inside.nssl.noaa.gov/mrms/past-code-updates/',
 'Stage IV': 'https://www.emc.ncep.noaa.gov/mmb/research/stage4.FAQ.html',
 'PRISM': 'https://www.prism.oregonstate.edu/calendar/',
 'PRISM grids': 'https://prism.oregonstate.edu/data/',
 'GFS': 'https://www.nco.ncep.noaa.gov/pmb/products/gfs/',
 'Cosgrove': 'https://doi.org/10.1029/2002JD003118',
}

class Slide:
    def __init__(self, title, subtitle='', section='PRODUCT GUIDE', refs=(), note=''):
        self.s = prs.slides.add_slide(prs.slide_layouts[6])
        self.s.background.fill.solid()
        self.s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        self.im = Image.new('RGB', (1600, 900), '#'+BG)
        self.d = ImageDraw.Draw(self.im)
        self.box(0,0,W,.1,TEAL)
        self.text(.5,.25,12,.28,section,11,TEAL,True)
        self.text(.5,.75,12.3,.7,title,32,INK,True)
        if subtitle: self.text(.5,1.48,12.3,.65,subtitle,18,GRAY)
        self.text(.5,7.12,11.7,.2,'HYDRO OPS  /  FORCING PRODUCTS  /  DESIGN SNAPSHOT: 19 SEP 2026',9,GRAY)
        self.text(12.35,7.05,.5,.3,f'{len(slides)+1:02}',12,TEAL,True)
        fullnote = note + '\n\nProject references: docs/forcing_production_workflow.md; docs/nrt_gfs_operations.md; docs/nwm_time_conventions.md.\n'
        fullnote += '\n'.join(f'{r}: {SOURCES[r]}' for r in refs)
        self.s.notes_slide.notes_text_frame.text = fullnote
        notes.append((title,fullnote))
        slides.append(self)

    def box(self,x,y,w,h,color=WHITE,rounded=False):
        sh=self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(color); sh.line.fill.background()
        xy=(round(x*SCALE),round(y*SCALE),round((x+w)*SCALE),round((y+h)*SCALE))
        if rounded:self.d.rounded_rectangle(xy,radius=14,fill='#'+color)
        else:self.d.rectangle(xy,fill='#'+color)
        return sh

    def text(self,x,y,w,h,value,size=22,color=INK,bold=False):
        # Explicit wrapping keeps the PowerPoint and preview line breaks identical.
        font=ImageFont.truetype(BOLD if bold else FONT,round(size*SCALE/72))
        lines=[]
        for para in str(value).split('\n'):
            line=''
            for word in para.split():
                candidate=(line+' '+word).strip()
                if self.d.textlength(candidate,font=font)>w*SCALE and line:
                    lines.append(line);line=word
                else:line=candidate
            lines.append(line)
        line_h=size/72*1.17
        if len(lines)*line_h > h+.025:
            raise ValueError(f'OVERFLOW slide {len(slides)}: {value!r}: {len(lines)*line_h:.2f} > {h}')
        sh=self.s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=sh.text_frame;tf.clear();tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold
            p.font.color.rgb=RGBColor.from_string(color)
            p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(size*1.17)
            self.d.text((round(x*SCALE),round((y+i*line_h)*SCALE)),line,font=font,fill='#'+color,anchor='lt')
        return sh

    def card(self,x,y,w,h,title,body='',color=TEAL):
        self.box(x,y,w,h,WHITE,True);self.box(x,y,.06,h,color)
        self.text(x+.2,y+.2,w-.4,.65,title,21,color,True)
        offset=.7 if h<2 else 1
        if body:self.text(x+.2,y+offset,w-.4,h-offset-.12,body,18 if h<2 else 20)

    def arrow(self,x,y,w=.45,color=TEAL):
        sh=self.s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x),Inches(y),Inches(w), Inches(.25))
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background()
        self.d.polygon([(x*SCALE,(y+.075)*SCALE),((x+w-.16)*SCALE,(y+.075)*SCALE),((x+w-.16)*SCALE,y*SCALE),((x+w)*SCALE,(y+.125)*SCALE),((x+w-.16)*SCALE,(y+.25)*SCALE),((x+w-.16)*SCALE,(y+.175)*SCALE),(x*SCALE,(y+.175)*SCALE)],fill='#'+color)

    def banner(self,value,color=TEAL,y=6.23):
        self.box(.5,y,12.3,.62,color,True);self.text(.72,y+.13,11.9,.42,value,20,WHITE,True)

def table(s, headers, rows, widths, y=2.15, row_h=1.1):
    x=.5
    for name,w in zip(headers,widths):
        s.box(x,y,w,.55,NAVY);s.text(x+.12,y+.12,w-.24,.3,name,16,WHITE,True);x+=w
    for i,row in enumerate(rows):
        x=.5
        for j,(value,w) in enumerate(zip(row,widths)):
            s.box(x,y+.55+i*row_h,w,row_h,WHITE if i%2==0 else LIGHT)
            s.text(x+.12,y+.74+i*row_h,w-.24,row_h-.2,value,19,TEAL if j==0 else INK,j==0);x+=w

# 1
s=Slide('Informed forcing for CONUS-wide NWM', 'A user guide to the data—not just the files', section='FORCING SUB-PROJECT')
for x,big,small,c in [(0.6,'1 km','NWM target grid',TEAL),(4.85,'1 hour','Time resolution',BLUE),(9.1,'2 streams','NRT + Retro',TEAL)]:
    s.box(x,2.65,3.65,2.4,WHITE,True);s.text(x+.25,3,3.2,.85,big,35,c,True);s.text(x+.25,4.2,3.2,.5,small,21)
s.banner('Best available evidence • Physically consistent fields • Traceable revisions')

# 2
s=Slide('Better informed—not simply higher resolution','Four design goals shape every production choice.',refs=('NLDAS forcing',))
for x,y,t,b in [(0.5,2.3,'Use observations','Prefer credible measurements; retain model coverage.'),(6.8,2.3,'Respect terrain','Adjust meteorology to NWM elevation.'),(0.5,4.15,'Keep consistency','Temperature, pressure, humidity and longwave stay coupled.'),(6.8,4.15,'Explain changes','Record sources, quality checks and revisions.')]:
    s.card(x,y,6,1.65,t,b)
s.banner('A finer grid is not proof of greater accuracy.')

#3
s=Slide('Eight fields, one model-ready product','Hourly meteorology on the NWM grid; daily files for efficient access.')
items=[('RAINRATE','Precipitation','kg m⁻² s⁻¹'),('T2D','2-m temperature','K'),('Q2D','2-m specific humidity','kg kg⁻¹'),('PSFC','Surface pressure','Pa'),('SWDOWN','Downward shortwave','W m⁻²'),('LWDOWN','Downward longwave','W m⁻²'),('U2D','10-m eastward wind','m s⁻¹'),('V2D','10-m northward wind','m s⁻¹')]
for i,(key,label,unit) in enumerate(items):
    x=.5+(i%4)*3.12;y=2.25+(i//4)*1.85
    s.box(x,y,2.95,1.65,WHITE,True);s.text(x+.17,y+.18,2.6,.4,key,23,TEAL,True);s.text(x+.17,y+.7,2.6,.65,label,18);s.text(x+.17,y+1.28,2.6,.25,unit,14,GRAY)
s.banner('Daily file = 24 hourly records. It does not mean daily-resolution data.')

#4
s=Slide('One production chain, two published streams','Native source grids map directly to NWM—no common 4-km intermediate grid.')
for x,t,b,c in [(0.5,'Acquire','Sources + revisions',BLUE),(3.65,'Build baseline','Select • remap • adjust',TEAL),(6.8,'Constrain','PRISM + consistency',TEAL),(9.95,'Publish','Audit • provenance',NAVY)]:
    s.card(x,2.8,2.85,2.2,t,b,c)
    if x<9:s.arrow(x+2.88,3.65,.25)
s.text(.7,5.38,5.9,.6,'Terrain • masks • cached weights',20,GRAY)
s.box(8,5.4,2.1,.6,BLUE,True);s.text(8.2,5.52,1.7,.4,'NRT',22,WHITE,True)
s.box(10.4,5.4,2.1,.6,TEAL,True);s.text(10.6,5.52,1.7,.4,'Retro',22,WHITE,True)
s.banner('Baseline is shared working data; NRT and Retro remain separate products.')

#5
s=Slide('Observation-informed inputs','Nominal source grids and update timing—not delivery guarantees.',refs=('MRMS','MRMS versions','Stage IV','PRISM','PRISM grids'),note='MRMS current operational tables indicate approximately 20-minute Pass 1 and 60-minute Pass 2 latency. Older versions used approximately 60 and 120 minutes. Stage-IV availability and revisions vary by RFC and accumulation period; no universal latency is asserted. PRISM uses the 4-km archive in this project, although other resolutions exist.')
table(s,['Source','Native grid / time','Availability / revisions','Main contribution'],[
['MRMS','~1 km / hourly QPE','Pass 1 ~20 min\nPass 2 ~60 min*','Storm timing + radar/gauge detail'],
['Stage-IV','~4 km / 1 h + 6 h','Hours; RFC-dependent\nLater revisions','Regional radar/gauge QC'],
['PRISM','4 km / daily','First ~1 day\nStable ~6 months','Rain totals + Tmin/Tmax']],[1.6,3,3.7,4],row_h=1.1)
s.text(.6,6.18,12,.6,'*Current MRMS latency. Older versions: ~1 / 2 hours. Actual availability varies.',16,GRAY)

#6
s=Slide('Complete meteorology across time and space','Complementary roles—not a “reanalysis always wins” rule.',refs=('NLDAS','NLDAS forcing','HRRR','GFS'))
table(s,['Source','Native grid / time','Availability / cycles','Role in this product'],[
['NLDAS-2','0.125° / hourly','~4-day lag','Consistent historical backbone'],
['HRRR','3 km / hourly','Hourly cycles\nPublication lag varies','Recent hourly analysis; forecast rain'],
['GFS','0.25° / hourly leads','00 / 06 / 12 / 18 UTC\nLeads 1–6; backup ≤12','Required northern NRT coverage']],[1.6,3,3.7,4],row_h=1.1)
s.banner('GFS is a short-forecast gap filler—not an hourly analysis or Retro source.')

#7
s=Slide('The same date becomes better informed','Illustrative data age—not a guaranteed product-delivery timeline.',refs=('PRISM','NLDAS','MRMS'))
stages=[('Hours','MRMS / Stage-IV\nHRRR + GFS'),('~1 day','Early PRISM\nconstraints'),('~4 days','NLDAS-2 replaces\nHRRR / GFS'),('~6 months','Stable PRISM\nRetro eligibility')]
for i,(title,body) in enumerate(stages):
    x=.5+i*3.15;s.card(x,2.6,2.8,2.5,title,body,TEAL if i==3 else BLUE)
    if i<3:s.arrow(x+2.87,3.55,.25)
s.box(.65,5.45,8.9,.48,BLUE,True);s.text(.85,5.51,8.5,.35,'NRT: mutable as observations and revisions arrive',18,WHITE)
s.box(10,5.45,2.7,.48,TEAL,True);s.text(10.15,5.51,2.4,.35,'Retro: stable',18,WHITE)
s.banner('Check actual completeness and revisions—not age alone.')

#8
s=Slide('Precipitation: select locally, constrain daily','Quality and coverage can change the preferred source from cell to cell.')
for i,(t,b) in enumerate([('Eligible sources','MRMS passes\nStage-IV\nNLDAS-2 → HRRR'),('Hourly composite','Quality + coverage\nRegional exceptions\nPreserve valid zeros'),('PRISM constraint','Adjust daily amount\nRetain hourly timing\nGuard dry/wet cases')]):
    s.card(.5+i*4.2,2.4,3.85,3.15,t,b)
    if i<2:s.arrow(4.42+i*4.2,3.7,.22)
s.banner('GFS fills unsupported northern rain only; it never overwrites valid rain.')

#9
s=Slide('PRISM changes the amount—not the storm clock','Schematic example: hourly pattern × 1.5 gives the PRISM daily total.',refs=('PRISM',),note='Illustration only, not observed data. Production uses PRISM 12–12 UTC accumulation windows, then regroups output into calendar-day files. Real reconciliation includes validity, dry/wet, scaling and domain checks, so not every cell follows an unconstrained ratio.')
vals=[0,0,1,2,4,2,1,0,0,1,1,0]
for panel,(x,f,c,title) in enumerate([(0.7,1,BLUE,'Hourly composite: 12 mm'),(7,1.5,TEAL,'PRISM-constrained: 18 mm')]):
    s.text(x,2.35,5.6,.5,title,23,c,True)
    for i,v in enumerate(vals):s.box(x+i*.42,5.5-v*f*.33,.3,max(.015,v*f*.33),c)
    s.box(x,5.5,5.15,.025,GRAY);s.text(x,5.7,5.4,.35,'Time within PRISM day →',17,GRAY)
s.arrow(6.02,3.8,.5)
s.banner('PRISM 12–12 windows are internal. Published files remain 00–23 UTC.')

#10
s=Slide('Terrain corrections stay physically coupled','Elevation-aware downscaling follows the Cosgrove framework.',refs=('Cosgrove',))
for x,t,b in [(0.5,'Temperature','Fixed lapse-rate\nelevation adjustment'),(3.65,'Pressure','Hydrostatic profile\nusing temperature'),(6.8,'Humidity','Interpolate RH\nRecompute specific humidity'),(9.95,'Longwave','Adjust effective\nemission temperature')]:
    s.card(x,2.5,2.85,2.7,t,b)
    if x<9:s.arrow(x+2.89,3.5,.22)
s.text(.7,5.55,11.9,.5,'PRISM temperature revisions also update humidity and longwave.',21,INK,True)
s.banner('Rotate wind vectors correctly. Preserve nighttime shortwave behavior.')

#11
s=Slide('The northern HRRR gap must not stop NWM','GFS completion is required for HRRR-selected NRT hours.',refs=('GFS',),note='Diagram is conceptual, not a geographic map. Actual coverage uses cached native-HRRR geometry and the versioned static envelope. Do not infer a latitude boundary from the rectangles. NLDAS-2 selection is per hour. GFS uses one meteorological bundle/cycle and elevation adjustments; valid precipitation is preserved. Missing required GFS fails the affected update, retaining previously accepted files.')
s.box(.7,2.5,6,3.25,LIGHT,True);s.box(.9,2.7,5.6,.8,GOLD);s.text(1.1,2.89,5.2,.5,'GFS: northern coverage',25,NAVY,True)
s.box(.9,3.55,5.6,1.98,BLUE);s.text(1.1,4.12,5.2,.7,'HRRR: native coverage',26,WHITE,True)
s.text(.8,5.92,6,.3,'Conceptual footprint—not a geographic map',14,GRAY)
s.card(7.1,2.5,5.55,1.5,'NLDAS-2 available?','Use NLDAS-2; no GFS needed.')
s.card(7.1,4.3,5.55,1.5,'Otherwise: HRRR + GFS','Missing required GFS → preserve old file.',BLUE)
s.banner('Keep active NWM cells covered; avoid unnecessary offshore extrapolation.')

#12
s=Slide('NRT and Retro answer different questions','Keep both: timeliness can be evaluated against the later stable product.')
for x,title,body,c in [(.5,'NRT','What could operations know?\n\nRecent, revisable sources\nEarly PRISM when eligible\nNLDAS replacement as it arrives',BLUE),(6.8,'Retro','What is our stable reconstruction?\n\nNLDAS-2 backbone\nBest eligible precipitation\nStable PRISM constraints',TEAL)]:
    s.card(x,2.25,6,3.75,title,body,c)
s.banner('Retro does not overwrite NRT. Record the version used in each simulation.')

#13
s=Slide('NRT: refresh → replace only what changed','Dependency fingerprints let unchanged baselines and windows be reused.')
for i,(t,b) in enumerate([('Refresh inputs','Check completeness\nand source revisions'),('Resolve changes','NLDAS arrival?\nPRISM revision?'),('Rebuild + audit','Reuse valid work\nReplace atomically')]):
    s.card(.5+i*4.2,2.4,3.85,2.8,t,b,BLUE)
    if i<2:s.arrow(4.42+i*4.2,3.45,.22)
s.text(.7,5.55,11.8,.5,'Required GFS acquisition is part of the recent-hour production path.',21,INK,True)
s.banner('A failed update must not replace accepted data with incomplete output.',BLUE)

#14
s=Slide('Scheduled opportunities ≠ current-hour delivery','Repository schedule template • not yet installed on the inspected host',note='Source: cron/hydro_ops.crontab and bin/update_nwm_forcing.py cycle_window, checked 2026-09-19. Six-hourly target scan 10 days with 14-day source repair lookback; daily 200 days; recent GFS worker tail seven days. Both NRT lanes end at UTC today minus two days. Monthly retro scans 45 target days ending today minus 183 days and still requires stable PRISM acceptance. These are scan windows, not mandatory rebuild counts. Faster event-driven/current-day delivery and daily retro eligibility scans are proposed, not activated.')
table(s,['Lane','UTC schedule','Current scan / behavior'],[
['NRT','02 / 14 / 20','10 target days; skip unchanged'],
['Daily NRT','08 each day','200-day deeper revision check'],
['Retro','10 on the 18th','45-day eligibility window; stable only']],[2.6,3,6.7],row_h=.95)
s.text(.7,5.92,11.9,.55,'NRT endpoint: UTC today − 2 days. No partial-current-day delivery.',20,RED,True)
s.text(.7,6.58,11.9,.3,'Status report template: every 2 hours • 7-day recent GFS-aware processing tail',15,GRAY)

#15
s=Slide('Retro: stable constraints, reproducible history','Source availability changes across eras; the common product contract does not.')
eras=[('1979–1980','NLDAS-2\nMonthly PRISM'),('1981–2001','NLDAS-2\nDaily PRISM'),('2002 onward','Add Stage-IV\nDaily PRISM'),('Modern era','Add MRMS / HRRR\nwhere eligible')]
for i,(t,b) in enumerate(eras):
    x=.5+i*3.15;s.card(x,2.6,2.8,2.6,t,b)
    if i<3:s.arrow(x+2.87,3.6,.25)
s.text(.7,5.55,11.9,.52,'Stable PRISM → final audit → safe baseline cleanup',22,INK,True)
s.banner('Coverage is still being built. Check the inventory—not just the intended era.')
s.s.notes_slide.notes_text_frame.text += '\n1979-01-01 is a partial 11-hour forcing day because NLDAS begins at 13 UTC. It is not a full-day model start. Cleanup is deferred in the active 2003–2020-10-13 campaign. Stable PRISM can change under a major source dataset release; preserve version provenance. HRRR anomaly refinement is not validated/enabled as a retrospective default.'

#16
s=Slide('One daily file; two clocks to understand','Storage dates and physical accumulation intervals are intentionally separate.')
s.text(.7,2.35,3.1,.5,'Files on disk',24,TEAL,True)
s.box(4,2.27,4.9,.7,TEAL,True);s.text(4.2,2.45,4.5,.45,'Day D: 00 … 23 UTC',23,WHITE,True)
s.box(9.1,2.27,3.5,.7,BLUE,True);s.text(9.3,2.45,3.1,.45,'Day D+1: 00 …',22,WHITE)
s.text(.7,3.75,3.1,.85,'24-hour NWM run',24,INK,True)
s.box(4.2,3.7,6.3,.7,NAVY,True);s.text(4.4,3.88,5.9,.45,'Reads D 01 … D+1 00',23,WHITE,True)
s.text(.7,5.05,3.1,.7,'PRISM rain day',24,INK,True)
s.box(4,5,8.6,.7,LIGHT,True);s.text(4.2,5.18,8.2,.45,'12–12 UTC window → recombine to calendar files',21)
s.banner('For a full UTC-day simulation, also supply hour 00 from the next file.')

#17
s=Slide('Quality safeguards users should know','A complete file is necessary—but not evidence of perfect accuracy.')
for x,y,t,b in [(.5,2.3,'CNRFC precipitation','Since July 2020: use trusted 6-hour totals, not striped hourly timing.'),(6.8,2.3,'Coverage envelope','Fill required active cells; control offshore extrapolation.'),(.5,4.2,'Physics + publication','Check coupled fields, timestamps, source provenance and file integrity.'),(6.8,4.2,'Known limits','No universal “best source”; independent validation remains future work.')]:
    s.card(x,y,6,1.75,t,b)
s.banner('Fine spacing + stable status + model-read success ≠ proven accuracy.')

#18
s=Slide('Find, inspect, then use','Consumers need the stream, valid-hour coverage and publication provenance.')
s.box(.5,2.3,12.3,1.1,NAVY,True);s.text(.75,2.57,11.8,.75,'forcing/outputs/conus/{nrt,retro}/YYYY/MM/\nYYYYMMDD.LDASIN_DOMAIN1',22,WHITE,True)
for i,(t,b) in enumerate([('Choose stream','NRT for operational replay\nRetro for stable history'),('Check status','python bin/\nreport_forcing_status.py'),('Check interval','24 records per day\nNext-day 00 for NWM')]):
    s.card(.5+i*4.2,3.8,3.85,2.1,t,b)
s.banner('Save source/version metadata with your experiment. NRT files can change.')
s.s.notes_slide.notes_text_frame.text += '\nRun from repository root in the hydro-ops environment: python bin/report_forcing_status.py. JSON: python bin/report_forcing_status.py --format json --output forcing/status/forcing-status.json. Status scans filenames/metadata, not a full scientific audit. Daily files deliberately have no .nc suffix. Baseline is an intermediate, not an alternative public stream.'

#19
s=Slide('What is ready—and what is not yet promised','User expectations as of 19 September 2026')
s.card(.5,2.4,6,3.5,'Implemented + tested','Required GFS northern fallback\nNLDAS arrival / PRISM revisions\nSeparate NRT and Retro products\nFailure-safe publication',TEAL)
s.card(6.8,2.4,6,3.5,'Still gated','Live cron coordination\nConsistent sub-hour delivery\nPartial-current-day publication\nIndependent accuracy validation',GOLD)
s.banner('Use the latest inventory for availability; use provenance for reproducibility.')

#20
s=Slide('Source guide & further reading','Official provider details and project methods are linked in slide notes.',section='REFERENCE / OPTIONAL APPENDIX')
for x,y,t,b in [(.5,2.25,'NOAA / NASA','MRMS • Stage-IV • HRRR • GFS\nNLDAS-2 forcing documentation'),(6.8,2.25,'Oregon State PRISM','Daily update calendar\nGrid definitions and revisions'),(.5,4.2,'Project documentation','Forcing production workflow\nNRT GFS operations + time conventions'),(6.8,4.2,'Acceptance evidence','NRT revision and reliability reports\nCoverage / quality limitations')]:
    s.card(x,y,6,1.75,t,b)
s.banner('Provider specifications describe inputs. Check local availability separately.')
s.s.notes_slide.notes_text_frame.text += '\n'+'\n'.join(f'{k}: {v}' for k,v in SOURCES.items())

path=OUT/'hydro_ops_forcing_user_guide.pptx'
prs.save(path)
for i,s in enumerate(slides,1):s.im.save(OUT/f'slide_{i:02}.png')
slides[0].im.save(OUT/'hydro_ops_forcing_user_guide_preview.pdf',save_all=True,append_images=[s.im for s in slides[1:]],resolution=120)
thumb=Image.new('RGB',(1600, math.ceil(len(slides)/4)*245),'#DCE5E9')
for i,s in enumerate(slides):
    im=s.im.resize((384,216));thumb.paste(im,(8+(i%4)*400,8+(i//4)*245))
thumb.save(OUT/'contact_sheet.png')
with (OUT/'speaker_notes.md').open('w') as f:
    for i,(title,note) in enumerate(notes,1):f.write(f'## {i}. {title}\n\n{slides[i-1].s.notes_slide.notes_text_frame.text}\n\n')
print(f'Saved {len(slides)} slides: {path}')
